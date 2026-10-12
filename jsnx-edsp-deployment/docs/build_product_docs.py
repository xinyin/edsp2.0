"""Rebuild the implementation-based product documentation with the bundled Python runtime."""
from pathlib import Path
import re
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).parent / '产品文档20261009'
OUT.mkdir(exist_ok=True)
BACK = ROOT / 'jsnx-edsp-platform-backend'
SCHEMA = BACK / 'platform-services/jsnx-edsp-data-admin/src/main/resources/db/schema.sql'
FONT = '/System/Library/Fonts/STHeiti Medium.ttc'

def diagram():
    w,h=1800,1600
    im=Image.new('RGB',(w,h),'white'); dr=ImageDraw.Draw(im)
    svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}"><rect width="100%" height="100%" fill="white"/>']
    def text(x,y,s,size=24):
        dr.text((x,y),s,font=ImageFont.truetype(FONT,size),fill='#172033')
        svg.append(f'<text x="{x}" y="{y+size}" font-family="PingFang SC, Microsoft YaHei, sans-serif" font-size="{size}" fill="#172033">{escape(s)}</text>')
    def box(x,y,width,title,lines):
        dr.rounded_rectangle((x,y,x+width,y+125),radius=8,fill='#f6f7f9',outline='#536477',width=2)
        svg.append(f'<rect x="{x}" y="{y}" width="{width}" height="125" rx="8" fill="#f6f7f9" stroke="#536477" stroke-width="2"/>')
        text(x+16,y+12,title,25)
        for i,s in enumerate(lines): text(x+16,y+53+i*29,s,20)
    def arrow(x1,y1,x2,y2,label='',dashed=False):
        dr.line((x1,y1,x2,y2),fill='#546879',width=3)
        import math
        a=math.atan2(y2-y1,x2-x1)
        pts=[(x2,y2),(x2-12*math.cos(a-.5),y2-12*math.sin(a-.5)),(x2-12*math.cos(a+.5),y2-12*math.sin(a+.5))]
        dr.polygon(pts,fill='#546879')
        svg.append(f'<path d="M{x1} {y1} L{x2} {y2}" stroke="#546879" stroke-width="3" fill="none"'+(' stroke-dasharray="8 6"' if dashed else '')+'/>')
        svg.append('<polygon points="'+' '.join(f'{x},{y}' for x,y in pts)+'" fill="#546879"/>')
        if label: text((x1+x2)/2+12,(y1+y2)/2-(30 if y1==y2 else 8),label,18)
    text(55,28,'外部数据平台 2.0 完整调用链路',36)
    text(55,82,'当前实现基线 2026-10-09   实线：业务请求与返回   独立分区：控制与异步观测',22)
    text(55,140,'一  线上业务调用与 XML 返回',28)
    box(55,205,430,'01  行内调用系统',['Service XML：service_id / service_sn / request','认证信息与调用系统 requester_id'])
    box(685,205,430,'02  模拟 ESB 请求队列',['请求入队与 requestReply 关联','本地为内存队列；生产对接行内 ESB'])
    box(1315,205,430,'03  ESB Adapter  8084',['Service XML → JSON，解析 service_id','向 Admin 解析已发布服务路径'])
    arrow(485,267,685,267);arrow(1115,267,1315,267)
    box(1315,440,430,'04  Gateway  8080',['校验凭证；向 Admin 查询 API 授权','本地路由匹配 code + method + path'])
    arrow(1530,330,1530,440,'JSON')
    box(685,440,430,'05  独立数据源应用',['校验应用归属与 X-Route-Version','读取本 JVM 的不可变 PlanSnapshot'])
    arrow(1315,502,1115,502)
    box(55,440,430,'06  请求转换与安全处理',['标准字段 → 供应商字段','透传 / 枚举 / 日期 / 加密 / 签名'])
    arrow(685,502,485,502)
    box(55,675,430,'07  实际上游接口',['按发布计划 URL 和超时调用','当前通用执行器使用 POST JSON'])
    arrow(270,565,270,675)
    box(685,675,430,'08  数据应用响应标准化',['提取自定义码值路径 → 统一返回码','供应商字段 → 行内字段；保留同一 Plan'])
    arrow(485,737,685,737,'上游响应')
    box(1315,675,430,'09  返回调用系统',['平台标准 JSON → Adapter → Service XML','ESB 响应队列 → 行内调用系统'])
    arrow(1115,737,1315,737)
    text(55,842,'失败分支：解析失败 / 授权拒绝 / 路由缺失 / 版本不一致 / 上游错误，返回对应错误并关联 Trace ID。',21)
    text(55,912,'二  配置发布与热更新控制链',28)
    box(55,973,430,'管理页面 → Admin  8081',['资源接入 → 标准服务配置 → 发布','测试可选；上线须配置校验通过'])
    box(685,973,430,'MySQL 持久化',['草稿 api_interface 及配置子表','线上 published_route / plan snapshot'])
    box(1315,973,430,'Gateway 与数据应用刷新',['Gateway 主动通知 + 5 秒定时拉取','数据应用通知刷新；缺失/版本不符重载'])
    arrow(485,1035,685,1035);arrow(1115,1035,1315,1035)
    text(55,1140,'下线删除线上路由但保留草稿；恢复旧发布快照与发布新草稿是两个独立操作。',22)
    text(55,1210,'三  异步日志与计量链',28)
    box(55,1270,430,'Gateway / 数据应用日志事件',['携带 Trace ID、状态与请求响应','Gateway 使用有界队列；失败可丢弃'])
    box(685,1270,430,'Observability  8083',['本地 HTTP 模拟 Kafka 事件传输','解析节点、调用结果并计算计量金额'])
    box(1315,1270,430,'MySQL → 运营页面',['usage_record / api_call_node_log','demo_async_event → 查询与统计'])
    arrow(485,1332,685,1332);arrow(1115,1332,1315,1332)
    text(55,1440,'实现边界：部分节点耗时为演示值；数据应用事件线程池仍需有界化；丢失日志不能作为可靠结算账。',21)
    text(55,1480,'控制面依赖：Adapter 解析、Gateway 授权及冷启动计划加载依赖 Admin。当前尚非完全离线自治。',21)
    im.save(OUT/'完整调用链路图.png');svg.append('</svg>');(OUT/'完整调用链路图.svg').write_text('\n'.join(svg))
    (OUT/'完整调用链路图.mmd').write_text('''sequenceDiagram
    participant C as 行内调用系统
    participant Q as 模拟ESB队列
    participant E as ESB Adapter
    participant G as Gateway
    participant A as Admin
    participant D as 独立数据源应用
    participant U as 上游供应商
    participant O as Observability
    participant M as MySQL
    C->>Q: Service XML与认证信息
    Q->>E: 消费请求并关联service_sn
    E->>A: service_id解析已发布路径
    A->>M: published_route关联data_source
    M-->>A: 路由记录
    A-->>E: 服务路径
    E->>G: JSON + API编码 + 调用系统
    G->>A: 授权有效期和日配额检查
    A-->>G: 授权结果
    alt 未授权或路由不可用
        G-->>E: 401 / 403 / 404 / 503
    else 路由可用
        G->>G: 匹配内存路由快照
        G->>D: published入口 + X-Route-Version
        opt 本地计划缺失或版本不一致
            D->>A: 拉取已发布计划与校验归属
            A->>M: release_plan_snapshot
            M-->>A: 持久化快照
            A-->>D: 不可变PlanSnapshot
        end
        D->>D: 请求映射和安全处理
        D->>U: POST JSON与配置超时
        U-->>D: 上游报文或HTTP错误
        D->>D: 返回码提取与响应标准化
        D-->>G: 统一JSON响应
        D-->>O: 异步调用事件
    end
    G-->>E: JSON响应
    G-->>O: 异步Gateway事件
    E->>E: 标准JSON封装为Service XML
    E-->>Q: XML响应入队
    Q-->>C: 按service_sn返回XML
    O->>M: 事件与节点日志和计量记录
    Note over D,O: Kafka事件保持结构化JSON；部分节点耗时为演示值
''',encoding='utf-8')

def new(title,intro):
    d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.75)
    for name,size in [('Normal',11),('Title',23),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
        st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts
        for attr in list(rf.attrib):
            if 'Theme' in attr: del rf.attrib[attr]
        rf.set(qn('w:eastAsia'),'Heiti SC')
        for border in st.element.xpath('.//w:pBdr'): border.getparent().remove(border)
        st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.15
    d.add_paragraph(title,'Title');d.add_paragraph('外部数据平台 2.0　代码基线 2026年10月9日')
    d.add_paragraph(intro)
    f=s.footer.paragraphs[0];f.alignment=2
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    return d
def p(d,t):d.add_paragraph(t)
def h(d,t):d.add_heading(t,1)
def sub(d,t):d.add_heading(t,2)
def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1, cols=len(headers));t.autofit=False
    widths=widths or [7/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for j,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if j==0:pr.append(OxmlElement('w:tblHeader'))
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i]);c.vertical_alignment=1
            tc=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');mar.append(e)
            tc.append(mar)
            if j==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');tc.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.05
                for r in para.runs:r.font.size=Pt(9);r.bold=j==0
    d.add_paragraph().paragraph_format.space_after=Pt(2)
def save(d,name):d.save(OUT/(name+'.docx'))

modules=[
('运营总览','全平台接口、资源、调用与计量概览；图表悬停查看统计信息','指标源于已落库记录，异步日志丢失会影响统计完整性'),
('资源目录','发现独立供应商开放的数据源与 API，接入并创建资源草稿','资源导入不代表已发布或已授权'),
('资源方管理','维护提供商编码、名称、合同、联系人、接入日期及账号权限资料','账号权限资料不等同于完整生产 IAM'),
('上游能力接入','维护资源 API 的地址、方法、请求响应样例和超时；实际调用上游验证','草稿可保存；只有 READY 资源可用于标准服务开发'),
('标准服务开发','引用资源 API，定义对外接口编码、路径、标准请求、映射、返回码、安全策略','请求与响应独立配置，无映射时该方向透传'),
('服务发布','配置校验、直接上线、历史快照、下线、恢复线上版本和归档','链路测试可选；上线与下线均保留历史'),
('API 访问授权','维护机构与调用系统，对接口进行单项或批量授权，配置环境、有效期和日配额','凭证校验与 API 授权不同；两项均通过才能调用'),
('计量管理','维护接口单价、命中计量条件，查询明细、金额和结果分布','现有异步日志可丢失，不是金融级最终结算账'),
('链路追踪','按调用记录和 Trace ID 查看成功失败、节点入出参和状态','全链路节点测量仍有模拟值，见详设边界'),
('运行监控','调用量、成功失败、耗时趋势、节点概况与告警展示','不是主机监控或全功能 APM')]

def product():
    d=new('外部数据平台产品功能介绍','本产品用于将多个外部资源提供商的接口接入、封装为行内标准服务，并向获授权的行内系统提供调用能力。本文面向业务负责人、产品经理和接入团队，说明功能边界、业务流程和交付能力。')
    h(d,'1 产品定位与服务对象')
    p(d,'平台统一管理资源接口与标准接口，降低新增接口及字段变更时的重复开发工作。平台维护对外标准契约，数据源应用在各自独立 JVM 中执行请求映射、签名加密和响应标准化。配置支持发布后动态生效；新增 Java 扩展算法仍需部署代码。')
    table(d,['角色','工作职责'],[('资源接入人员','维护资源方与上游 API，核对契约并联调'),('服务开发人员','定义行内标准契约，配置映射、安全策略和返回码'),('发布与授权人员','发布、下线、恢复版本，为机构和系统授权'),('运营运维人员','核对调用结果、计量记录和异常链路'),('行内调用系统','按标准接口文档发送 Service XML 请求至 ESB，并消费 Service XML 响应')],[1.4,5.6])
    h(d,'2 功能模块全景');table(d,['功能模块','能力','使用边界'],modules,[1.25,3.2,2.55])
    h(d,'3 资源与标准服务的关系')
    p(d,'一个资源提供商包含多个数据源，一个数据源包含多个资源 API。标准服务通过 resource_api_id 关联资源 API；该字段没有唯一约束，因此同一资源 API 可被不同标准服务复用。每个标准服务具有唯一 interface_code，调用报文中的 apiCode 使用这个标准服务编码。')
    p(d,'资源 API 保存供应商请求与返回的基准样例；标准服务持有契约副本、映射、安全策略和发布版本。上游变更后，通过同步资源契约显式更新草稿。已发布快照不随资源样例编辑而变化。')
    h(d,'4 配置化开发能力')
    table(d,['能力','操作与结果'],[('请求转换','行内标准字段 → 供应商请求字段，支持默认值、必填及转换规则'),('响应转换','供应商返回字段 → 行内标准字段；与请求方向分别配置'),('透传','某个方向无映射时保留原始对象结构；统一响应外层仍由平台生成'),('字段规则','直接映射、常量、大小写、枚举、日期、安全转换'),('统一码值','配置供应商码值路径，例如 status.code；映射到标准码与 HIT / MISS / ERROR'),('安全策略','选择注册算法与密钥引用；支持国密和 Java SPI 扩展'),('超时继承','默认采用数据源超时，可为单个接口覆盖'),('复杂报文','对象、嵌套路径与数组路径能力按当前解析器验证；不承诺任意 JSONPath 或任意 XML 数组自动识别')],[1.2,5.8])
    h(d,'5 开发发布与授权流程')
    p(d,'资源建草稿 → 保存上游配置 → 可选上游联调 → 基于资源创建标准服务 → 配置并保存 → 可选草稿链路测试 → 配置校验与上线 → 调用系统授权 → 可选线上入口验收。审批记录和接口保留为扩展，当前主流程可直接上线。')
    p(d,'开发与线上运行彼此分开：线上 V2 可以同时存在 V3 草稿。下线只停止线上请求，草稿继续开发。恢复 V2 与发布 V3 分别使用旧快照与新配置；保存不会自动恢复生产流量。')
    h(d,'6 端到端业务调用');d.add_picture(str(OUT/'完整调用链路图.png'),width=Inches(6.2))
    h(d,'7 当前能力与交付边界')
    p(d,'当前本地版本使用 MySQL、独立 Spring Boot 服务与 Vue 3 页面，无容器依赖。ESB 采用内存队列模拟，异步事件采用 HTTP 模拟 Kafka。Nacos、Sentinel 和 Kafka 有接入配置，默认未开启；Router 仅预留契约。')
    p(d,'商用交付前须补齐生产认证授权、行内 ESB 适配、真实节点测量、数据脱敏与权限校验、事件队列有界化、数据库变更管理以及容量验收。产品支持配置发布不代表代码已经完成银行生产验收。日千万次约等于平均 116 次每秒，不能代替峰值压测、故障演练和容量规划。')
    save(d,'01 产品功能介绍')

def manual():
    d=new('外部数据平台使用手册','本手册面向首次接入和开发人员，以资源 API 接入、标准服务封装、上线授权和排障为主线。操作前应准备供应商接口文档、脱敏请求响应样例、运行地址与调用系统资料。示例编码仅用于说明，不代表数据库已有对应配置。')
    h(d,'1 启动与登录')
    p(d,'部署目录为 jsnx-edsp-deployment。确认 JDK 17 和本地 MySQL 可用后执行 ./bin/start-demo.sh。平台页面为 http://127.0.0.1:5173/，供应商独立页面为 http://127.0.0.1:5190/。当前开发登录账号 admin，密码 admin123；此固定登录机制仅供本地验证。')
    p(d,'停止使用 ./bin/stop-demo.sh。后端代码更新可使用 ./bin/reload-integration-stack.sh 重载联调栈，该脚本不重置数据。不要在正常使用时执行 reset-test-data.sh。修改前端样式后 Vite 会热更新；静态部署需要重新构建并替换 dist。')
    h(d,'2 页面导航与操作约定')
    p(d,'侧边菜单按全局视图、数据服务和运行管理分组，可收起分组及整个侧栏。接口资源目录按数据源折叠，选择 API 后在右侧编辑。离开有未保存修改的页面时，选择继续编辑或明确放弃。保存只保存开发态配置，上线才切换线上版本。')
    h(d,'3 准备上游供应商与资源')
    for title,body in [
    ('3.1 建立资源提供商','进入资源方管理，新增或编辑提供商编码、名称、联系方式、合同及接入属性。提供商编码应稳定，避免把接口编码当成提供商编码。'),
    ('3.2 建立数据源','建立数据源编码、基础 URL 和默认超时，并登记所属独立应用 application_name 与 runtime_url。runtime_url 是数据应用地址，如 http://127.0.0.1:8191；endpoint_url 是供应商基础地址，两者不可混用。实例数字段不负责启动进程。'),
    ('3.3 创建上游 API','进入上游能力接入，选择数据源创建草稿，填写供应商路径、请求样例和返回样例。默认继承数据源配置；仅需特殊地址或超时的接口使用覆盖配置。也可从资源目录选择供应商能力后接入并创建草稿。'),
    ('3.4 保存和验证','点击保存上游配置，接口成为 READY，才可在标准服务开发中选择。点击验证上游接口时必须按供应商要求输入参数。若供应商要求 query_key，传 company_code 不会自动替代；标准映射属于下一阶段。'),
    ('3.5 核对真实数据','供应商模拟服务支持按数据库中的企业记录查询，校验失败、命中和未命中应分别验证。新增 API 定义不会自动生成真实企业记录。')]:sub(d,title);p(d,body)
    h(d,'4 创建标准服务')
    p(d,'进入标准服务开发，选择已保存的资源 API 创建服务，填写唯一服务编码、服务名称、对外路径和 HTTP 方法。行内系统与 ESB 交互采用 Service XML；Adapter 到 Gateway 及数据应用内部执行链使用 JSON。其他供应商方法需确认对应运行实现，不应仅凭页面可选就认定支持。')
    p(d,'明确行内系统调用的标准请求与响应契约。供应商字段本身已符合标准时，保留请求方向和响应方向的无映射状态即可。透传表示业务对象不逐项转换，最终仍有 code、message、resultStatus、traceId、data 的平台统一外层。')
    table(d,['方向','来源','目标','示例'],[('请求','行内标准请求','供应商请求','companyCode → query_key'),('响应','供应商响应','行内标准响应','result.company_name → companyName'),('响应枚举','供应商原始值','行内标准值','sex_code 的 1 → M，2 → F')],[.6,1.4,1.4,3.6])
    sub(d,'4.1 编辑字段映射');p(d,'粘贴标准请求或供应商响应 JSON 样例，使用样例导入字段生成映射行，再确认来源路径与目标路径。请求和响应可独立选择透传或转换。新增行需要填写完整路径；启用映射后只输出配置的目标字段，未配置字段不会自动追加。')
    sub(d,'4.2 配置枚举和日期');p(d,'枚举抽屉左侧填写实际来源值，右侧填写目标标准值；保存抽屉后还需保存字段标准化。未知枚举可配置默认值。日期填写来源格式与目标格式，例如 yyyy-MM-dd → yyyyMMdd，错误格式会导致运行异常。')
    sub(d,'4.3 返回码');p(d,'填写供应商码值取值路径，如 result.code 或 status.code，再添加原始码与平台码。HIT 表示命中业务数据，MISS 表示正常未命中，ERROR 表示业务或供应商错误。HTTP 200 不保证 HIT。参与计量控制是否生成计费金额，允许重试是策略标记，不能据此认定已经执行自动重试。')
    sub(d,'4.4 签名与加密');p(d,'从后台注册算法目录选择算法，填写受控密钥引用和必要参数，按供应商签名规范确认字段顺序、字符集和编码。密钥引用可采用 config://、env://、sys://。新增自定义算法需要实现 Java SPI 并部署扩展 JAR；页面配置不能创建任意 Java 算法。')
    sub(d,'4.5 保存和测试');p(d,'保存后可执行配置校验，检查映射及样例契约。草稿链路测试为可选测试功能，会真实调用上游，但不经过 ESB 和 Gateway。修改上游契约后点击同步资源契约，再核对标准请求和映射；同步不改变线上版本。')
    h(d,'5 上线下线与版本')
    table(d,['操作','结果'],[('首次反复保存','current_version 为 1，published_version 为空'),('首次上线','线上 V1，生成发布快照与路由'),('修改已上线服务并保存','形成 V2 草稿，线上仍 V1'),('继续保存同一草稿','V2 不因保存次数增加'),('下线','停止线上流量，仍可编辑草稿'),('恢复线上 V1','恢复旧快照，不发布 V2 草稿'),('上线草稿 V2','通过配置校验后发布 V2'),('删除或归档','未发布草稿可删除；有历史的接口先下线再归档，日志与发布记录保留')],[1.65,5.35])
    p(d,'服务发布页面查看当前草稿与已发布版本。没有链路测试记录不阻止上线，但静态配置校验失败仍应修复。历史审批功能不是当前主流程的前置条件。')
    h(d,'6 调用系统与批量授权')
    p(d,'进入 API 访问授权，先建立调用系统及法人机构信息，再选择接口设置环境、起止日期和日配额。批量操作前使用数据源和搜索筛选，核对已选集合及数量，再保存。确认调用请求的 X-App-Id 与 client_code 完全一致，environment 与授权环境一致。')
    p(d,'配置调用系统不等于完成凭证发放。当前 Gateway 开发凭证为配置项 edsp.auth.bootstrap-api-key；正式环境需对接独立的调用方凭证体系。下线后授权记录保留，但无法调用；恢复线上后有效授权可继续使用。')
    h(d,'7 Postman 完整链路验收')
    p(d,'先完成上线、运行地址登记及授权，再使用 POST http://127.0.0.1:8084/esb/queue/request。每次调用使用不同 Trace ID，避免 usage_record 唯一键将多次调用合并。')
    table(d,['请求头','填写内容'],[('Content-Type','application/xml'),('X-App-Id','已授权的 client_code'),('X-Api-Key','当前环境配置的调用凭证'),('X-Trace-Id','POSTMAN-20261009-0001，每次递增或使用 UUID')],[1.6,5.4])
    p(d,'请求示例，替换 apiCode 为已上线的标准服务编码，body 使用该服务标准请求字段：')
    p(d,'<Service>\n  <Service_Header>\n    <service_sn>TRACE-001</service_sn>\n    <service_id>COMPANY_QUERY_SERVICE</service_id>\n    <requester_id>LOAN_APPROVAL</requester_id>\n    <service_time>20261012140000</service_time>\n    <version_id>01</version_id>\n  </Service_Header>\n  <Service_Body>\n    <request><companyCode>DEMO-001</companyCode></request>\n  </Service_Body>\n</Service>')
    p(d,'预期响应结构：{"code":"EXT-SUCCESS","message":"查询成功","resultStatus":"HIT","traceId":"...","data":{...}}。实际标准码和 data 以已发布配置为准；未命中应检查 MISS，不能仅检查 HTTP 状态。')
    h(d,'8 调用追踪与计量核对');p(d,'在链路追踪按 Trace ID 查记录，点击节点查看输入输出。依次确认 ESB 解析参数、Gateway 目标、实际供应商请求和响应转换结果。日志异步入库，短时查不到可稍后刷新；传输故障允许丢日志，不保证每笔都有完整节点。计量管理按接口规则显示单价和金额，最终结算需可靠数据渠道。')
    h(d,'9 常见问题')
    table(d,['现象','检查顺序'],[('保存按钮不可用','确认当前资源与编辑内容已加载，检查必填字段和网络错误；已下线不应禁止保存'),('供应商 400 缺少字段','先看实际上游请求，再对比已保存资源契约；确认请求映射方向'),('Gateway 403','调用系统、接口、环境、有效期、启用状态和日配额'),('Gateway 503 无运行地址','data_source.runtime_url 与实际应用端口、application_name 归属'),('ESB 等待超时','查看 Adapter、Gateway、上游耗时；外层超时应覆盖内层与处理余量'),('路由不存在','确认上线及 published_route；检查 code、HTTP 方法和路径一致性'),('上线提示缺字段','同步上游契约，检查映射后的实际字段及默认值；测试不替代静态校验'),('页面已变但接口仍旧','确认重载了对应后端 JAR；页面热更新不会重启 Java'),('新配置未生效','确认保存后已上线；查看 route_version、发布快照和数据应用选中版本')],[1.65,5.35])
    save(d,'02 使用手册')

TABLES={
'data_provider':('资源提供商','提供商档案及合同和接入属性'), 'provider_account':('资源方账号','账号角色权限资料；不存认证密码'),
'data_source':('数据源','供应商基础连接与数据应用归属'), 'resource_api':('上游资源接口','供应商原始契约和接入状态'),
'published_route':('已发布路由','线上服务路径、版本和目标数据源'), 'api_interface':('标准服务草稿','行内标准服务与资源副本、版本和配置'),
'field_mapping':('字段映射','请求与响应逐字段规则'), 'interface_validation_record':('验证记录','配置摘要与可选验证时间'),
'api_security_policy':('接口安全策略','算法模式及密钥引用'), 'response_code_mapping':('返回码映射','供应商码到标准结果和计量标记'),
'vendor_mock_config':('内置模拟配置','历史内置模拟上游；独立供应商不依赖此表'), 'release_record':('发布记录','版本发布结果与操作者'),
'release_plan_snapshot':('发布执行计划','持久化不可变执行快照'), 'release_approval':('发布审批','保留审批扩展流程'),
'interface_lifecycle_record':('生命周期记录','上线、下线、恢复及归档操作审计'), 'api_client':('调用系统','系统编码与法人机构'),
'api_authorization':('API 授权','系统与接口在指定环境的授权及配额'), 'billing_rule':('计量规则','接口单价与计量配置'),
'usage_record':('调用计量明细','按 Trace ID 归集调用结果与金额'), 'api_call_node_log':('节点日志','按 Trace ID 保存节点状态和入出参'),
'demo_async_event':('本地异步事件','HTTP 模拟 Kafka 接收的事件原文'),
'vendor_data_source':('供应商数据源','独立供应商的资源目录'), 'vendor_api_definition':('供应商 API 定义','模拟接口契约与返回模板'),
'vendor_enterprise_record':('企业样例数据','真实查询命中和未命中的演示基础记录')}
FIELDS=dict(x.split('=',1) for x in '''id=记录主键
provider_code=资源提供商唯一编码
provider_name=资源提供商名称
credit_code=统一社会信用代码
provider_type=提供商类型
contact_name=联系人姓名
access_date=接入日期
message_encoding=通讯字符编码
charge_required=是否收费
encryption_required=是否要求加密
login_required=是否要求登录
contact_phone=联系人电话
contact_email=联系人邮箱
service_phone=服务电话
contract_no=合同编号
contract_start_date=合同开始日期
contract_end_date=合同结束日期
security_level=安全等级标记
settlement_mode=结算周期标记
office_address=办公地址
description=说明
status=记录业务状态
created_at=记录创建时间
updated_at=记录更新时间
provider_id=所属提供商主键
username=账号名称
display_name=展示姓名
mobile=手机号码
email=电子邮箱
role_code=角色编码
permissions=权限项 JSON
last_login_at=最后登录时间
source_code=数据源唯一编码
source_name=数据源名称
application_name=独立数据源应用名称与归属校验值
runtime_url=数据应用运行地址或生产 VIP
protocol_type=协议类型标记
endpoint_url=供应商基础地址
default_timeout_ms=数据源默认超时 毫秒
environment=环境编码 与授权环境匹配
instance_count=实例数配置元数据 不创建实际进程
config_version=数据源配置修订号
data_source_id=所属数据源主键
resource_api_code=上游资源 API 编码
resource_api_name=上游资源 API 名称
http_method=HTTP 方法标记
upstream_url=供应商完整 URL
upstream_path=供应商相对路径
config_scope=INHERIT 继承或 OVERRIDE 覆盖
request_example=请求样例 JSON
response_example=响应样例 JSON
timeout_ms=接口超时 毫秒 为空时继承数据源
current_version=当前配置版本 资源为修订号 标准服务为候选发布号
verified_at=最近上游验证时间
verified_checksum=验证配置摘要
interface_id=标准服务主键
interface_code=标准服务唯一编码 对应 apiCode
service_path=平台标准服务 URL 路径
route_version=当前线上路由版本
published_at=路由发布时间
resource_api_id=来源资源 API 关联
interface_name=标准服务名称
dispatch_mode=调度模式 当前 DIRECT
router_service_code=预留 Router 服务标识
request_schema=请求结构定义 JSON
response_schema=响应结构定义 JSON
upstream_request_example=供应商请求基准副本
upstream_response_example=供应商响应基准副本
standard_response_example=行内标准响应样例
retry_count=重试次数配置 不代表当前链路已自动重试
vendor_code_path=供应商返回码取值路径
upstream_status=上游接入状态 DRAFT 或 READY
published_version=最近发布版本 下线后仍保留
mapping_direction=REQUEST 或 RESPONSE
source_path=来源字段路径
target_path=目标字段路径
transform_type=字段转换类型
transform_config=转换参数 JSON
default_value=字段缺失或规则默认值
required_flag=是否必填
sort_order=执行或展示顺序
validation_type=验证类型
config_checksum=验证绑定的配置摘要
validated_at=验证时间
enabled=功能启用开关
encrypt_mode=加密算法编码
signature_mode=签名算法编码
request_fields=参与安全处理字段
secret_reference=历史密钥引用字段 当前使用分离的两个引用
encryption_key_reference=加密密钥引用 不保存明文密钥
signing_key_reference=签名密钥引用 不保存明文密钥
signature_field=签名写入字段
vendor_code=供应商业务码
standard_code=平台标准业务码
result_status=HIT 命中 MISS 未命中 ERROR 错误
retryable=重试策略标记
billable=是否参与计量
mock_name=模拟配置名称
response_body=供应商返回模板 JSON
http_status=模拟 HTTP 状态
delay_ms=模拟延迟 毫秒
release_version=发布版本号
release_status=发布结果
checksum=配置或计划摘要
release_note=发布说明
released_by=发布操作者
released_at=发布时间
plan_json=完整发布执行计划 JSON
applicant=申请人
approval_node=审批节点
submitted_at=提交时间
reviewed_by=审批人
review_comment=审批意见
reviewed_at=审批时间
action=生命周期动作
operated_by=操作者
reason=操作原因
operated_at=操作时间
client_code=调用系统唯一编码 对应 X-App-Id
client_name=调用系统名称
institution_code=法人机构编码
institution_name=法人机构名称
owner_name=负责人
client_id=调用系统主键 可空时表示未关联已登记系统
daily_quota=每日调用配额
valid_from=授权生效日期
valid_to=授权截止日期
approved_by=授权经办人
billing_mode=计量模式标记
unit_price=调用单价
currency=币种
hit_billable=命中计量配置
miss_billable=未命中计量配置
effective_date=规则生效日期
trace_id=贯穿链路的调用标识
amount=本次调用计量金额
duration_ms=记录的调用耗时 毫秒
upstream_duration_ms=上游调用耗时 毫秒
gateway_duration_ms=网关耗时 当前部分路径为占位值
request_status=SUCCESS 或 FAILED
node_name=应用或节点名称
called_at=调用记录时间
node_key=节点标识
node_category=节点类别
node_status=节点结果状态
node_detail=节点说明
input_data=节点输入 JSON
output_data=节点输出 JSON
event_type=事件类型
payload=事件内容 JSON
consume_status=事件消费状态标记
base_path=供应商数据源基础路径
api_code=供应商 API 编码
api_name=供应商 API 名称
response_code=模拟业务返回码
version=供应商 API 修订版本
enterprise_code=企业内部查询编码
enterprise_name=企业名称
legal_person=法定代表人
sex_code=供应商性别码
risk_level=风险等级
registration_status=工商登记状态'''.splitlines())

def schema_tables(path):
    return [(m.group(1),m.group(2)) for m in re.finditer(r'CREATE TABLE IF NOT EXISTS (\w+)\s*\((.*?)\n\);',path.read_text(),re.S)]

def design():
    d=new('外部数据平台详细设计与数据库设计','本文面向架构、开发、测试和运维团队，覆盖当前代码的模块职责、调用契约、配置发布、运行链路与数据库字段。基线是本仓库实际实现；生产化目标与未实现项单独列明，避免将演示机制视为已完成的生产保障。')
    h(d,'1 架构约束与技术选型')
    table(d,['层次','当前实现'],[('前端','Vue 3 + TypeScript + Element Plus + Vite；前后端目录独立'),('后端','JDK 17、Spring Boot、Spring Cloud Gateway；Maven 多模块'),('数据','MySQL edsp_demo；供应商独立 vendor_demo'),('部署','独立 JVM 与普通脚本；不引入容器与 Spring Cloud LoadBalancer'),('配置与治理','Nacos 和 Sentinel 接入配置默认关闭；具体上线能力须按环境验证'),('事件','Gateway 有界异步日志；本地 HTTP 模拟 Kafka'),('运行能力','独立数据源应用本地 PlanSnapshot；集中 capability-runtime 为可选项')],[1.3,5.7])
    h(d,'2 独立部署模块与公共组件')
    table(d,['应用','端口','职责'],[('data-admin','8081','资源、标准服务配置、发布快照、授权、计量规则、草稿测试'),('data-gateway','8080','入口认证与授权、路由匹配、异步入口日志'),('esb-adapter','8084','Service XML 与内部 JSON 双向转换、模拟请求响应队列'),('data-observability','8083','接收事件、节点和计量落库、日志及监控查询'),('data-enterprise','8181','企业数据独立执行单元'),('data-credit','8182','征信数据独立执行单元'),('data-antifraud','8183','百融反欺诈独立执行单元'),('data-source-lhzx','8191','联合征信独立执行单元'),('capability-runtime','8091','可选复杂能力服务 普通链路不必经'),('vendor-backend','8090','独立上游模拟系统')],[1.65,.6,4.75])
    p(d,'工程模块名均带 jsnx-edsp- 前缀。公共组件包括 contracts、client-sdk、plan-core、transform-core、crypto-core、signature-core、response-core、log-core、platform-core；其中 platform-core 承载管理侧领域逻辑，数据应用通过 SDK 使用本地计划执行能力。现有目录仍为多个 Maven 组件，不能把此前讨论的单一 core 目录方案当成已经落地的结构。')
    p(d,'独立 JVM 提供进程隔离；Admin、MySQL、Gateway 仍为共享依赖，因此不等于全系统绝对故障隔离。公共 JAR 修复需要升级引用它的应用，使用接口和 SPI 可减少业务代码变动，但不能免除部署。')
    h(d,'3 完整线上调用链');d.add_picture(str(OUT/'完整调用链路图.png'),width=Inches(7))
    sub(d,'3.1 入口解析与路由');p(d,'Adapter 校验 Service XML，提取 service_id、service_sn、requester_id 与 Service_Body/request，转换为内部 JSON。service_id 作为平台标准服务编码解析已发布路径，service_sn 贯穿全链路。Adapter 再向 Gateway POST；Gateway 认证 X-App-Id 和 X-Api-Key，并调用 Admin 检查授权。路由过滤器同时校验 X-Api-Code、HTTP 方法和路径。')
    sub(d,'3.2 数据应用执行');p(d,'Gateway 将请求转发至 runtime_url + /internal/adapter/published，并附加 X-Route-Version。LocalPlanRuntime 检查应用归属，加载或复用对应版本的计划。单次调用捕获同一个 PlanSnapshot，依次执行请求映射、安全处理、上游 POST、返回码映射和响应标准化。在途请求继续使用已捕获的快照。')
    sub(d,'3.3 响应与异常');p(d,'平台内部响应外层包含 code、message、resultStatus、traceId、data。Adapter 将成功或失败结果统一封装为 Service XML，回显原请求头和请求体后写入 ESB 响应链路。Kafka 日志与事件仍保存结构化 JSON。业务 ERROR 与 HTTP 状态不能混用，监控应同时记录二者。')
    h(d,'4 控制面与数据面数据模型')
    p(d,'data_provider → data_source → resource_api → api_interface 构成业务资源层级。api_interface 的 resource_api_id 可空且非唯一，支持同一资源派生多个标准服务。field_mapping、response_code_mapping、api_security_policy 是标准服务子配置。')
    p(d,'api_interface 保存当前编辑配置；published_route 保存已发布的路径、方法和版本；release_plan_snapshot 保存不可变执行内容。路由目标地址由 data_source.runtime_url 关联获得，因此调整数据源运行地址将影响该数据源全部路由，须作为运维变更管理。')
    h(d,'5 发布版本与不停机生效')
    p(d,'保存草稿将 current_version 置为 COALESCE(published_version,0)+1；反复保存同一草稿不继续增号。资源 API 的 current_version 是资源修订号，语义不同。草稿测试记录绑定摘要且保存后失效，但不作为上线前置条件。')
    p(d,'直接发布在事务内执行配置校验并写入发布记录、快照和路由。提交后通知数据应用拉取计划，再通知 Gateway 刷新。Gateway 启动时拉取，之后默认每 5 秒全量刷新，并以 AtomicReference 替换本地快照。通知失败时轮询兜底，所以不同实例不是分布式瞬时原子切换。')
    p(d,'数据应用内存计划缺失时通过 Admin 从持久化快照加载，无需重启；收到版本不一致请求时刷新一次，仍不一致则拒绝执行。进程崩溃则需要进程重启，重启后首次调用恢复计划；Admin 不可用时冷启动无法恢复。Gateway 刷新失败保留上一有效快照，但冷启动且拉取失败时无可用路由。')
    p(d,'下线删除 published_route，撤销活动计划并保留发布快照。OFFLINE 状态下仍可编辑保存；恢复旧线上版本使用发布快照，新草稿单独发布。发布与下线操作应记录操作者和原因。当前单状态字段承载运行态和开发态的部分语义，后续可迁移为两个独立状态字段。')
    h(d,'6 转换与安全设计')
    p(d,'PlanMappingSupport 按方向执行映射；空集合复制原对象。启用规则后输出新对象，仅写入已配置目标路径。默认值在读取值缺失或空串时使用；枚举、常量、大小写、日期、加密和签名基于 transform_type 与 transform_config。数组及嵌套能力须以 FieldTransformSupport 的路径语法为准。')
    p(d,'LocalPlanRuntime 从 vendor_code_path 提取原始码并匹配 response_code_mapping。未匹配默认 ERROR。当前仅 HIT 分支执行响应映射，MISS / ERROR 的标准数据为空。安全算法由 SecurityAlgorithmRegistry 注册表和 Java SPI 提供；密钥引用在执行应用环境中解析，发布计划不应存明文密钥。')
    h(d,'7 授权与计量设计')
    p(d,'api_authorization 使用 client_id、interface_id、environment 联合唯一键，每条记录包含有效期、状态和日配额。当前 Gateway 每次向 Admin 查询授权，存在同步控制面依赖。当前配额基于已落库记录统计，异步延迟和丢失可能使配额不精确；生产应采用原子配额策略。')
    p(d,'观测服务 consumeCall 根据接口查询单价，以事件 billable 标记计算 amount，并以 trace_id 唯一键写 usage_record。当前此路径未完整执行规则生效日期和所有计量模式；日志丢失、重复 Trace ID、单价变更时点会影响结算准确性。正式账务应独立建立可靠投递、幂等、规则版本及对账机制。')
    h(d,'8 日志与故障边界')
    p(d,'Gateway 日志通过 log-core 有界队列发送，故障允许丢弃并计数。数据应用 DemoEventClient 当前采用单线程无界任务队列，虽然异步异常不回传业务，但持续故障下可能积压内存；生产必须改成有界队列、限长报文、超时和丢弃计数。')
    p(d,'demo_async_event 保存事件原文，api_call_node_log 使用 trace_id + node_key 唯一键，usage_record 使用 trace_id 唯一键。日志可含敏感数据，必须统一字段脱敏、访问控制和保留周期。节点重复覆盖和异步乱序需明确合并策略。')
    p(d,'StandardDataApplicationExecutor 和 Adapter 诊断中部分 ESB、Gateway、计量耗时固定为 1ms，部分成功状态由模板组装；不能据此开展真实性能定位。生产应由每个实际执行节点记录起止时间、状态及入出参摘要，再按 Trace ID 聚合。')
    h(d,'9 管理接口目录')
    controller=BACK/'platform-services/jsnx-edsp-data-admin/src/main/java/com/jsnx/edsp/web/ApiController.java'
    apis=re.findall(r'@(Get|Post|Put|Delete)Mapping\("([^\"]+)"\)',controller.read_text())
    table(d,['方法','路径'],[(m.upper(),'/api'+path) for m,path in apis],[.65,6.35])
    p(d,'内部管理端点必须限制网络访问并接入服务身份校验。当前固定开发登录与共享 bootstrap-api-key 不属于商用身份体系。接口表仅展示当前路由；请求 DTO 以 ApiModels 和 Controller 为准。')
    h(d,'10 数据库总体设计')
    p(d,'平台库 edsp_demo 与独立供应商库 vendor_demo 均使用 utf8mb4。下列字段字典来自仓库 schema.sql 的 CREATE TABLE 定义，包含全部字段、类型、约束及业务用途；实际环境若经过历史 ALTER，应比对 SHOW CREATE TABLE。没有读取或导出业务行数据。')
    platform=schema_tables(SCHEMA);vendor=schema_tables(ROOT/'jsnx-edsp-vendor-backend/src/main/resources/db/schema.sql')
    table(d,['数据库','表数量','用途'],[('edsp_demo',len(platform),'平台配置、发布、授权、调用与计量'),('vendor_demo',len(vendor),'独立供应商模拟资源与企业数据')],[1.4,.65,4.95])
    table(d,['表','中文名称','用途'],[(name,*TABLES[name]) for name,_ in platform+vendor],[2,1.2,3.8])
    sub(d,'10.1 关联与约束');p(d,'配置子表关联 api_interface，多数配置子表设置 ON DELETE CASCADE。发布记录、授权、计量明细外键不使用级联删除，业务采用归档保留历史。published_route 未声明外键，依赖发布事务维持一致性；节点日志按 trace_id 逻辑关联，不具备数据库外键。')
    sub(d,'10.2 数据保留与性能');p(d,'高频日志表应按时间组织归档和清理，不建议长期与控制面配置共用无限增长的表。索引包括接口与调用时间、调用状态与时间、Trace ID 与节点顺序。当前多数 DATETIME 未声明毫秒精度，即使 NOW(3) 写入也不能保证保留毫秒；需要迁移为 DATETIME(3) 才能满足原始时间记录要求。')
    h(d,'11 数据库字段字典')
    for n,(name,body) in enumerate(platform+vendor,1):
        d.add_page_break();sub(d,f'11.{n} {TABLES[name][0]}')
        p(d,('vendor_demo.' if name.startswith('vendor_') and name!='vendor_mock_config' else 'edsp_demo.')+name+'　'+TABLES[name][1])
        rows=[];constraints=[]
        for raw in body.splitlines():
            line=raw.strip().rstrip(',')
            if not line:continue
            if re.match(r'(CONSTRAINT|UNIQUE|INDEX|PRIMARY|KEY)\b',line):constraints.append(line);continue
            match=re.match(r'(\w+)\s+([A-Z]+(?:\([^)]*\))?)(.*)',line)
            if not match:raise ValueError(line)
            field,typ,rest=match.groups()
            if field not in FIELDS:raise ValueError('Missing field meaning: '+field)
            flags=[]
            if 'PRIMARY KEY' in rest:flags.append('PK')
            if 'AUTO_INCREMENT' in rest:flags.append('自增')
            if 'UNIQUE' in rest:flags.append('唯一')
            flags.append('非空' if 'NOT NULL' in rest or 'PRIMARY KEY' in rest else '可空')
            default=re.search(r'DEFAULT\s+(\S+)',rest)
            if default:flags.append('默认 '+default.group(1))
            if 'ON UPDATE' in rest:flags.append('自动更新时间')
            rows.append((field,typ,'；'.join(flags),FIELDS[field]))
        table(d,['字段','类型','约束与默认值','用途'],rows,[1.75,1.05,1.8,2.4])
        if constraints:
            p(d,'索引与外键')
            for constraint in constraints:p(d,constraint)
    h(d,'12 新数据源开发与部署流程')
    p(d,'1 在 data-applications 创建独立 Spring Boot 模块，使用 JDK 17，引入 contracts、client-sdk 及实际必要组件；配置唯一应用名和端口。2 注册数据源，填写与进程一致的 application_name 和 runtime_url。3 实现供应商特殊算法或协议的扩展，通用字段转换采用配置。4 接入资源 API 并建立标准服务。5 单元、契约、草稿和入口测试后上线并授权。6 以普通 JVM 部署，多实例通过行内 VIP 对外。')
    p(d,'配置变更通过发布快照动态生效；Java 算法变更通过构建、版本化发布与滚动更新完成。核心组件缺陷应按受影响依赖范围升级、灰度及回滚，不能承诺只更新一处就让已加载旧 JAR 的 50 个进程自动修复。')
    h(d,'13 验收与生产化清单')
    table(d,['验收项','判定标准'],[('契约转换','透传、嵌套、数组、必填、枚举、日期和错误输入分别有契约样例'),('热发布','并发请求验证新旧版本不混用；刷新通知丢失后可收敛'),('生命周期','下线可开发；恢复旧版不带草稿；归档保留历史'),('授权','未授权、过期、错误环境、停用和超配额均被拒绝'),('故障隔离','供应商慢响应不拖垮其他应用；共享 Admin 故障影响明确'),('日志压力','长报文、队列满、观测不可用时进程内存有上界'),('真实追踪','每个节点真实起止时间与状态，日志去敏，失败链可追溯'),('数据治理','迁移脚本、备份恢复、归档周期及权限审计'),('容量','按峰值与慢上游并发压测，不以日均吞吐代替 SLA'),('安全','移除固定登录和共享示例凭证，接入 IAM 与密钥系统，内部接口鉴权')],[1.3,5.7])
    h(d,'14 实现定位')
    for f in ['platform-services/jsnx-edsp-data-admin/src/main/java/com/jsnx/edsp/web/ApiController.java','common-components/jsnx-edsp-platform-core/src/main/java/com/jsnx/edsp/service/InterfaceService.java','platform-services/jsnx-edsp-data-gateway/src/main/java/com/jsnx/edsp/gateway/PublishedRouteRegistry.java','common-components/jsnx-edsp-client-sdk/src/main/java/com/jsnx/edsp/sdk/LocalPlanRuntime.java','common-components/jsnx-edsp-client-sdk/src/main/java/com/jsnx/edsp/sdk/StandardDataApplicationExecutor.java','platform-services/jsnx-edsp-data-observability/src/main/java/com/jsnx/edsp/observability/DemoEventIngestionService.java']:p(d,f)
    save(d,'03 详细设计与数据库设计')

if __name__=='__main__':
    diagram();product();manual();design()
    print(OUT)
